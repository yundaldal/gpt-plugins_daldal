"""품의봇 자동 점검(하네스): 만들어진 기안문(HWPX)·엑셀을 case.json과 독립적으로 다시 확인한다.

사용법: python harness.py case.json out_dir   (out_dir/점검결과.md 작성, 실패가 있으면 종료코드 1)

점검: (1) 금액 일치 (2) 붙임 (3) 항목기호·날짜 (4) 엑셀 재계산 (5) 개인정보 (6) 원고 (7) HWPX 구조.
세금·한글 금액은 amount.py를 쓰지 않고 이 파일에서 따로 계산해 대조한다.
"""
import json
import glob
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from datetime import date
from fractions import Fraction

import openpyxl
from lxml import etree

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)
from assemble import fill_draft_amounts  # noqa: E402
from hwpx_check import extract_paragraphs  # noqa: E402
from readiness import load_config  # noqa: E402

HWPX_NAME = "기안문.hwpx"
XLSX_NAMES = {"allowance": "수당 지급 내역.xlsx", "purchase": "물품 구입 목록.xlsx"}
BUDGET_NAME = "예산연결 업로드양식 결과.xlsx"
BUDGET_FORM = os.path.join(BASE_DIR, "..", "resources", "기준", "양식", "예산연결 업로드양식.xlsx")
RATES_PATH = os.path.join(BASE_DIR, "..", "resources", "기준", "수당단가.json")
MARKS = "가나다라마바사아자차카타파하"
WEEK = "월화수목금토일"
DATE_LOOSE = re.compile(r"\d{4}\s*[.\-/]\s*\d{1,2}\s*[.\-/]\s*\d{1,2}(?!\d)")
DATE_STRICT = re.compile(r"\d{4}\. (?:[1-9]|1[0-2])\. (?:[1-9]|[12]\d|3[01])")
DATE_WEEKDAY = re.compile(r"(?:(\d{4})\. )?(\d{1,2})\. (\d{1,2})\.(?:\(([^)])\))?")
ACCOUNT_LIKE = re.compile(r"(?<![\d-])\d{3,6}-\d{2,6}-\d{2,8}(?:-\d{1,4})?(?![\d-])|(?<![\d,.])\d{10,16}(?![\d,])")
AMOUNT_TEXT = re.compile(r"금([\d,]+)원\(금([가-힣]+)원\)")
SAFE_IDENTIFIER = re.compile(r"(?:cell:[A-Z]{1,3}\d{1,5}|field:[a-z_][a-z0-9_.\[\]]*)")
TAX_OTHER_EXEMPTION = 125000
TAX_OTHER_RATE = Fraction(8, 100)
TAX_BUSINESS_RATE = Fraction(3, 100)
TAX_LOCAL_RATE = Fraction(10, 100)
TAX_FLOOR_UNIT = 10


def _norm(s):
    return re.sub(r"\s+", "", s or "")


def _ids(*identifiers):
    """Build a report detail from safe schema paths and spreadsheet cell IDs only."""
    return "ids:" + " ".join(identifiers)


def _safe_report_detail(detail, fallback):
    identifiers = list(dict.fromkeys(SAFE_IDENTIFIER.findall(detail or "")))
    if identifiers:
        return "확인 위치: " + ", ".join(identifiers)
    return f"확인 위치: {fallback}"


# ---------- 독립 계산 (amount.py와 별도 구현) ----------

def korean_amount(n):
    """4자리씩 끊어 읽는다. 공문 관례대로 십·백·천 앞의 '일'도 적는다(10만 -> 일십만)."""
    if n == 0:
        return "영"
    digit = dict(zip("0123456789", "영일이삼사오육칠팔구"))
    s = str(n)
    s = s.zfill(-(-len(s) // 4) * 4)
    groups = [s[i:i + 4] for i in range(0, len(s), 4)]
    out = ""
    for gi, g in enumerate(groups):
        unit_big = ["", "만", "억", "조"][len(groups) - 1 - gi]
        part = "".join(digit[ch] + u for ch, u in zip(g, ["천", "백", "십", ""]) if ch != "0")
        if part:
            out += part + unit_big
    return out


def floor10(x):
    """10원 미만 절사(Fraction 입력)."""
    x = Fraction(x)
    return int(x - x % TAX_FLOOR_UNIT)


def calc_tax(income_type, total, base):
    """Recalculate from an independent constant snapshot; verify current policy separately."""
    if income_type == "기타":
        inc = floor10(Fraction(total) * TAX_OTHER_RATE) if total > TAX_OTHER_EXEMPTION else 0
    elif income_type == "사업":
        inc = floor10(Fraction(base) * TAX_BUSINESS_RATE)
    else:
        raise ValueError("소득 구분 값이 지원되지 않습니다.")
    return inc, floor10(Fraction(inc) * TAX_LOCAL_RATE)


# ---------- 엑셀 읽기 ----------

def cell_value(ws, coord, depth=0):
    """숫자 또는 '=1+2', '=K5+L5', '=J5-M5' 형태의 직접 합산·참조식을 값으로 풀어 준다. 엑셀 계산 결과에는 의존하지 않는다."""
    v = ws[coord].value
    if isinstance(v, (int, float)):
        return int(v)
    if isinstance(v, str) and v.startswith("=") and depth < 3:
        body = v[1:]
        if re.fullmatch(r"\d+(\+\d+)*", body):
            return sum(int(t) for t in body.split("+"))
        m = re.fullmatch(r"([A-Z]+\d+)([+-])([A-Z]+\d+)", body)
        if m:
            a, b = cell_value(ws, m.group(1), depth + 1), cell_value(ws, m.group(3), depth + 1)
            return a + b if m.group(2) == "+" else a - b
    raise ValueError(f"{coord} 값을 읽을 수 없음: {v!r}")


def is_numeric_cell(value):
    """True only for a stored numeric cell, never for a formula string."""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def first_amount(text):
    m = re.match(r"\s*([\d,]+)원", text or "")
    if not m:
        raise ValueError(f"금액 문구가 아님: {text!r}")
    return int(m.group(1).replace(",", ""))


def total_source(case):
    """총액 출처 우선순위는 수당 > 경비 > 물품이며 C3 물품은 경비 총액에 더하지 않는다."""
    for key in ("allowance", "expense", "purchase"):
        if case.get(key):
            return key
    return None


def independent_totals(case, out_dir):
    """엑셀·입력에서 총액을 다시 구해 {출처: 총액}."""
    totals = {}
    al = case.get("allowance")
    if al:
        path = os.path.join(out_dir, XLSX_NAMES["allowance"])
        ws = openpyxl.load_workbook(path).active
        n = len(al["people"])
        totals["allowance"] = sum(cell_value(ws, f"J{5 + i}") for i in range(n))
    pu = case.get("purchase")
    if pu:
        totals["purchase"] = sum(int(it["qty"]) * int(it["unit_price"]) for it in pu["items"])
    ex = case.get("expense")
    if ex:
        t = 0
        for ln in ex["lines"]:
            amt = int(ln["unit_price"])
            for f in ln.get("factors") or []:
                amt *= int(f)
            t += amt
        totals["expense"] = t
    return totals


# ---------- 점검 1: 금액 ----------

def formula_amounts(texts):
    """'단가*수량=금액' 줄의 금액들."""
    return [int(m.group(1).replace(",", "")) for t in texts if "=" in t and re.search(r"[*×]", t)
            for m in [re.match(r"\s*([\d,]+)원", t.rsplit("=", 1)[1])] if m]


def check_amounts(case, out_dir, texts):
    res = []
    comp = case.get("computed") or {}
    total = comp.get("total")
    res.append(("computed.total이 계산돼 있음", total is not None and not comp.get("questions"), _ids("field:computed.total")))
    if total is None:
        return res
    found = AMOUNT_TEXT.findall("\n".join(texts))
    bad = [i for i, (n, k) in enumerate(found)
           if int(n.replace(",", "")) != total or k != korean_amount(total)]
    if found:
        res.append(("기안문 금액 = computed.total, 한글 금액 일치", not bad,
                    _ids(*(f"field:draft.amount[{i}]" for i in bad)) if bad else ""))
    else:
        # 협의회비처럼 총액 문구 없이 산출식 줄만 적는 기안문은 산출식 금액의 합으로 대조한다
        parts = formula_amounts(texts)
        res.append(("기안문에 총액 문구가 없으면 산출식 금액 합 = computed.total",
                    bool(case.get("expense")) and bool(parts) and sum(parts) == total,
                    _ids("field:draft.expense.amounts")))
    tl = comp.get("total_line", "")
    res.append(("computed.total_line 한글 금액 일치", tl == f"금{total:,}원(금{korean_amount(total)}원)", _ids("field:computed.total_line")))
    try:
        totals = independent_totals(case, out_dir)
    except (OSError, ValueError, KeyError):
        res.append(("엑셀·입력에서 총액 재계산", False, _ids("field:recalculation")))
        return res
    src = total_source(case)
    res.append(("엑셀(또는 입력) 재계산 합계 = computed.total", totals.get(src) == total,
                _ids(f"field:{src}.total", "field:computed.total")))
    return res


# ---------- 점검 2: 붙임 ----------

def parse_attachments(texts):
    """('붙임'으로 시작하는 줄부터) -> (번호 또는 None, 이름 줄) 목록, 시작 인덱스."""
    start = max((i for i, t in enumerate(texts) if t.startswith("붙임")), default=None)
    if start is None:
        return [], None
    items = []
    for t in texts[start:]:
        m = re.match(r"(?:붙임\s+)?(\d+)\.\s+(.*)$", t)
        if m:
            items.append((int(m.group(1)), m.group(2)))
        elif t.startswith("붙임"):
            items.append((None, re.sub(r"^붙임\s+", "", t)))
    return items, start


def check_attachments(case, texts, config):
    res = []
    items, start = parse_attachments(texts)
    res.append(("붙임 줄이 있음", start is not None, _ids("field:draft.attachments")))
    if start is None:
        return res
    nums = [n for n, _ in items]
    consecutive = nums == [None] if len(items) == 1 else nums == list(range(1, len(items) + 1))
    res.append(("붙임 번호가 1부터 연속", consecutive, _ids("field:draft.attachments.numbering")))
    res.append(("마지막 줄이 '끝.'으로 끝남", texts[-1].rstrip().endswith("끝."), _ids("field:draft.terminator")))
    res.append(("'끝.'은 마지막 줄에만", sum(t.count("끝.") for t in texts[start:]) == 1, _ids("field:draft.terminator")))
    names = [_norm(name) for _, name in items]
    for att in config["유형"][case["type"]]["붙임"]:
        if not att["required"]:
            continue
        keys = [att["key"]] + att.get("keywords", [])
        hit = any(_norm(k) in n for k in keys for n in names)
        res.append((f"필수 붙임 포함: {att['key']}", hit, _ids("field:draft.attachments")))
    return res


# ---------- 점검 3: 항목기호·날짜 ----------

def check_symbols_dates(case, texts):
    res = []
    start = max((i for i, t in enumerate(texts) if t.startswith("붙임")), default=len(texts))
    body = texts[:start]
    first = next((i for i, t in enumerate(body) if re.match(r"1\. ", t)), None)
    res.append(("본문 '1.' 시작 줄이 있음", first is not None, ""))
    if first is None:
        return res
    errors, top, mark_i = [], 0, 0
    for line_i, t in enumerate(body[first:], start=first):
        m = re.match(r"(\d+)\. ", t)
        if m:
            top += 1
            mark_i = 0
            if int(m.group(1)) != top:
                errors.append(f"field:draft.items[{line_i}]")
            continue
        m = re.match(r"([가-하])\. ", t)
        if m:
            if m.group(1) != MARKS[mark_i]:
                errors.append(f"field:draft.items[{line_i}]")
            mark_i += 1
        elif not t.startswith("- "):
            errors.append(f"field:draft.items[{line_i}]")
    res.append(("항목기호 순서(1. → 가. → -)", not errors, _ids(*errors) if errors else ""))

    bad, in_related = [], False
    date_texts = [t for t in texts if t.startswith("제목")] + body[first:]
    for text_i, t in enumerate(date_texts):
        is_title = t.startswith("제목")
        t_plain = re.sub(r"^(\d+|[가-하])\. ", "", t)
        if re.match(r"\d+\. ", t):
            in_related = t_plain.startswith("관련")
        for m in DATE_LOOSE.findall(t_plain):
            if not DATE_STRICT.fullmatch(m):
                bad.append(f"field:draft.date[{text_i}]")
        year = None
        for m in DATE_WEEKDAY.finditer(t_plain):
            y, mo, d, wd = m.groups()
            year = int(y) if y else year
            if year is None:
                bad.append(f"field:draft.date[{text_i}]")
                continue
            try:
                real = WEEK[date(year, int(mo), int(d)).weekday()]
            except ValueError:
                bad.append(f"field:draft.date[{text_i}]")
                continue
            if wd is None:
                if not (in_related or is_title):
                    bad.append(f"field:draft.date[{text_i}]")
            elif wd != real:
                bad.append(f"field:draft.date[{text_i}]")
    res.append(("날짜 형식 'YYYY. M. D.(요일)'과 요일 정확성", not bad, _ids(*bad) if bad else ""))
    return res


# ---------- 점검 4: 엑셀 재계산 ----------

def check_allowance_xlsx(case, out_dir):
    res = []
    path = os.path.join(out_dir, XLSX_NAMES["allowance"])
    if not os.path.exists(path):
        return [("수당 지급 내역.xlsx 존재", False, _ids("field:allowance.workbook"))]
    ws = openpyxl.load_workbook(path).active
    people = case["allowance"]["people"]
    n = len(people)
    errs, numeric_errs, sums = [], [], dict.fromkeys("JKLMN", 0)
    for i, p in enumerate(people):
        r = 5 + i
        try:
            g, h, t = (first_amount(ws[f"{c}{r}"].value) for c in "GHI")
            for col in "JKLMN":
                if not is_numeric_cell(ws[f"{col}{r}"].value):
                    numeric_errs.append(f"cell:{col}{r}")
            j = cell_value(ws, f"J{r}")
            inc, loc = calc_tax(p.get("income_type"), j, g + h)
            want = {"J": g + h + t, "K": inc, "L": loc, "M": inc + loc, "N": g + h + t - inc - loc}
            if ws[f"E{r}"].value != p.get("name"):
                errs.append(f"cell:E{r}")
            for col, w in want.items():
                got = cell_value(ws, f"{col}{r}")
                sums[col] += got
                if got != w:
                    errs.append(f"cell:{col}{r}")
        except (ValueError, KeyError):
            errs.append(f"field:allowance.people[{i}]")
    res.append(("엑셀 J~N(총액·세금·공제·실수령액) 독립 재계산 일치", not errs, _ids(*errs)))
    res.append(("수당 개인별 J~N 셀은 계산된 숫자", not numeric_errs,
                _ids(*numeric_errs)))
    tot = 5 + n
    last = tot - 1
    total_numeric = [c for c in "JKLMN" if not is_numeric_cell(ws[f"{c}{tot}"].value)]
    ok = ws[f"A{tot}"].value == "합계" and not total_numeric and all(
        ws[f"{c}{tot}"].value == sums[c] for c in "JKLMN")
    detail = "" if ok else _ids(*(f"cell:{c}{tot}" for c in set(total_numeric) | {c for c in "JKLMN" if ws[f"{c}{tot}"].value != sums[c]}))
    res.append(("합계 행(계산된 숫자)", ok, detail))
    comp = (case.get("computed") or {}).get("allowance")
    if comp:
        diff = [c for c, k in zip("JKLMN", ("total", "income_tax", "local_tax", "deduction", "net")) if sums[c] != comp[k]]
        res.append(("엑셀 열 합계 = computed.allowance", not diff,
                    _ids(*(f"cell:{c}{tot}" for c in diff)) if diff else ""))
    return res


def _hours(s):
    """인정시간이 있으면 그 값, 없으면 시작~종료 시각 차이(시간)."""
    if s.get("recognized_hours") not in (None, ""):
        return Fraction(str(s["recognized_hours"]))
    a, b = (int(x) for x in s["start"].split(":"))
    c, d = (int(x) for x in s["end"].split(":"))
    return Fraction((c * 60 + d) - (a * 60 + b), 60)


def _by_day(sessions):
    days = {}
    for s in sessions:
        days.setdefault(s["date"], []).append(_hours(s))
    return list(days.values())


def _round_half_up(x):
    return int(Fraction(x) + Fraction(1, 2))


def expected_instructor_pay(person, rates):
    """강사수당(G) 독립 계산: 강의일마다 기본+초과×(인정시간-1), 30분 미만 반액, 그룹은 별도."""
    days = _by_day(person["sessions"])
    cat = person["category"]
    if cat == "그룹":
        g = rates["그룹강사"]
        extra_n = max(0, int(person["group_size"]) - g["기본인원"])
        return (g["기본료"] + g["초과_1인당"] * extra_n) * len(days)
    r = rates["강사수당"][cat]
    half = rates["반액규칙"]
    pay = Fraction(0)
    for hs in days:
        h = sum(hs, Fraction(0))
        if h < Fraction(str(half["기준_시간"])):
            pay += r["기본"] * Fraction(str(half["비율"]))
        else:
            pay += r["기본"] + r["초과"] * max(Fraction(0), h - 1)
    return _round_half_up(pay)


def expected_manuscript_pay(person, rates):
    """원고료(H) 독립 계산: min(매수, 시간×한도)×단가, 재활용 수정률 30% 미만 0 / 70% 이하 50% / 초과 100%."""
    m = person.get("manuscript") or {}
    kind = m.get("kind") or "none"
    if kind == "none" or m.get("meets_spec") is False:
        return 0
    r = rates["원고료"][kind]
    pages = Fraction(str(m["pages"]))
    if kind == "도서":
        pay = min(r["단가"] * pages, r["지급한도액"])
    else:
        hours = sum((sum(hs, Fraction(0)) for hs in _by_day(person["sessions"])), Fraction(0))
        pay = r["단가"] * min(pages, hours * r["시간당한도"])
    ratio = m.get("revision_ratio")
    if ratio is not None:
        pct = Fraction(str(ratio))
        pay *= 0 if pct < 30 else Fraction(1, 2) if pct <= 70 else 1
    return _round_half_up(pay)


def expected_session_pay(kind, person, rates):
    """컨설팅: 회차마다 기본 + 3시간 넘는 회차가 있는 날 초과 1회. 위원회: 날마다 기본 + 하루 2시간 이상이면 초과 1회."""
    days = _by_day(person["sessions"])
    if kind == "consulting":
        r = rates["컨설팅"][person["category"]]
        limit = rates["컨설팅"]["기본시간"]
        n = sum(len(hs) for hs in days)
        return r["기본"] * n + r["초과"] * sum(1 for hs in days if any(h > limit for h in hs))
    c = rates["위원회"]
    return c["기본"] * len(days) + c["초과"] * sum(1 for hs in days if sum(hs, Fraction(0)) >= c["초과_기준시간"])


def check_allowance_rates(case, out_dir):
    """엑셀 G·H 첫 금액과 J를 기준/수당단가.json + case 입력만으로 다시 계산해 대조한다(amount.py 미사용)."""
    path = os.path.join(out_dir, XLSX_NAMES["allowance"])
    if not os.path.exists(path):
        return []
    with open(RATES_PATH, encoding="utf-8") as f:
        rates = json.load(f)
    ws = openpyxl.load_workbook(path).active
    kind = case["allowance"].get("kind", "instructor")
    errs = []
    for i, p in enumerate(case["allowance"]["people"]):
        r = 5 + i
        try:
            if kind == "instructor":
                g, h = expected_instructor_pay(p, rates), expected_manuscript_pay(p, rates)
            else:
                g, h = expected_session_pay(kind, p, rates), 0
            travel = int((p.get("travel") or {}).get("amount") or 0)
            got_g, got_h = first_amount(ws[f"G{r}"].value), first_amount(ws[f"H{r}"].value)
            got_j = cell_value(ws, f"J{r}")
        except (ValueError, KeyError, TypeError):
            errs.append(f"field:allowance.people[{i}]")
            continue
        for col, want, got in (("G", g, got_g), ("H", h, got_h), ("J", g + h + travel, got_j)):
            if want != got:
                errs.append(f"cell:{col}{r}")
    return [("엑셀 G·H·J를 단가표로 독립 재계산한 금액과 일치", not errs, _ids(*errs))]


def check_purchase_xlsx(case, out_dir):
    path = os.path.join(out_dir, XLSX_NAMES["purchase"])
    if not os.path.exists(path):
        return [("물품 구입 목록.xlsx 존재", False, _ids("field:purchase.workbook"))]
    ws = openpyxl.load_workbook(path).active
    items = case["purchase"]["items"]
    errs = []
    for i, it in enumerate(items):
        r = 4 + i
        if (ws[f"B{r}"].value, ws[f"D{r}"].value, ws[f"F{r}"].value) != (it["name"], it["qty"], it["unit_price"]):
            errs.extend(f"cell:{col}{r}" for col in "BDF")
        if ws[f"G{r}"].value != int(it["qty"]) * int(it["unit_price"]):
            errs.append(f"cell:G{r}")
    tot = 4 + len(items)
    expected_total = sum(int(it["qty"]) * int(it["unit_price"]) for it in items)
    sum_ok = ws[f"A{tot}"].value == "합계" and ws[f"G{tot}"].value == expected_total
    return [("물품 목록 행(품명·수량·단가·금액식) 일치", not errs, _ids(*errs)),
            ("물품 목록 합계 행(숫자 합계)", sum_ok, _ids(f"cell:G{tot}"))]


def expected_budget_rows(case):
    """B(물품)·C(일반 지출)의 예산연결 행 [(품명, 수량, 단가)]. build_budget_upload와 별개로 다시 계산한다."""
    if case.get("purchase"):
        return [(it["name"], it["qty"], it["unit_price"]) for it in case["purchase"]["items"]]
    rows = []
    for ln in case["expense"]["lines"]:
        qty = 1
        for f in ln.get("factors") or []:
            qty *= int(f)
        rows.append((ln["label"], qty, int(ln["unit_price"])))
    return rows


def check_budget_upload(case, out_dir):
    path = os.path.join(out_dir, BUDGET_NAME)
    has_budget_type = bool(case.get("purchase") or case.get("expense"))
    if not has_budget_type:
        return [("예산연결 업로드 파일 없음(수당 유형)", not os.path.exists(path), _ids("field:budget_upload"))]
    if not os.path.exists(path):
        return [("예산연결 업로드 파일 존재", False, _ids("field:budget_upload.workbook"))]
    ws = openpyxl.load_workbook(path).active
    form = openpyxl.load_workbook(BUDGET_FORM).active
    rows = expected_budget_rows(case)
    n = len(rows)
    errs = []
    if [c.value for c in ws[1][:18]] != [c.value for c in form[1][:18]]:
        errs.append("field:budget_upload.headers")
    data_rows = [r for r in range(2, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in range(1, 19))]
    if data_rows != list(range(2, 2 + n)):
        errs.append("field:budget_upload.rows")
    total = 0
    for i, (name, qty, price) in enumerate(rows):
        r = 2 + i
        want = {"A": i + 1, "B": name, "C": qty, "D": None, "E": price,
                "F": qty * price, "L": case["school_level"]}
        for c in range(1, 19):
            col = openpyxl.utils.get_column_letter(c)
            got = ws.cell(r, c).value
            exp = want.get(col)
            if got != exp:
                errs.append(f"cell:{col}{r}")
        if ws[f"D{r}"].value not in (None, ""):
            errs.append(f"cell:D{r}")
        c_val, e_val, f_val = ws[f"C{r}"].value, ws[f"E{r}"].value, ws[f"F{r}"].value
        if isinstance(c_val, int) and isinstance(e_val, int):
            total += c_val * e_val
        if f_val != qty * price:
            errs.append(f"cell:F{r}")
    expected_total = case["computed"]["total"] if case.get("computed", {}).get("total") is not None else None
    if expected_total is not None and total != int(expected_total):
        errs.extend(["field:computed.total"] + [f"cell:F{r}" for r in range(2, 2 + n)])
    return [("예산연결 업로드 파일(머리글·행·값·금액식·학교급·합계)", not errs, _ids(*errs))]


def check_expense(case, texts):
    """기안문의 '단가*수량=금액' 줄 산술 + 경비 입력 줄 금액."""
    res, bad = [], []
    for text_i, t in enumerate(texts):
        if "=" not in t or not re.search(r"[*×]", t):
            continue
        left, right = t.rsplit("=", 1)
        m = re.match(r"\s*([\d,]+)원", right)
        if not m:
            continue
        expr = left.partition(": ")[2] or left
        prod = 1
        for f in re.split(r"[*×]", expr):
            fm = re.match(r"\s*([\d,]+)", f)
            if not fm:
                prod = None
                break
            prod *= int(fm.group(1).replace(",", ""))
        if prod != int(m.group(1).replace(",", "")):
            bad.append(f"field:draft.items[{text_i}]")
    res.append(("기안문 산출식(단가×수량=금액) 산술", not bad, _ids(*bad)))
    errs = []
    for i, ln in enumerate(case["expense"]["lines"]):
        amt = int(ln["unit_price"])
        for f in ln.get("factors") or []:
            amt *= int(f)
        if ln.get("amount") not in (None, "") and int(ln["amount"]) != amt:
            errs.append(f"field:expense.lines[{i}].amount")
    res.append(("경비 입력 금액 = 단가×수량", not errs, _ids(*errs)))
    return res


# ---------- 점검 5: 개인정보 ----------

def output_texts(out_dir):
    """결과물(엑셀 모든 칸, HWPX 모든 XML)의 텍스트 {파일: 문자열}."""
    out = {}
    for path in sorted(glob.glob(os.path.join(out_dir, "*.xlsx"))):
        parts = []
        for ws in openpyxl.load_workbook(path):
            parts += [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        out[os.path.basename(path)] = "\n".join(parts)
    for path in sorted(glob.glob(os.path.join(out_dir, "*.hwpx"))):
        with zipfile.ZipFile(path) as zf:
            # 본문 문단과 문서 정보(content.hpf)만 본다. header.xml 등의 식별 번호는 계좌와 무관하다
            out[os.path.basename(path)] = "\n".join(
                extract_paragraphs(zf) + ["".join(etree.fromstring(zf.read("Contents/content.hpf")).itertext())])
    return out


def check_privacy(case, out_dir):
    """Check that account-shaped values in generated files came from this case.

    This narrow runtime check is not a general personal-data detector. Release archives
    are scanned separately by the local build helper before they are handed off.
    """
    case_text = _norm(json.dumps(case, ensure_ascii=False))
    unprovided = 0
    for text in output_texts(out_dir).values():
        for m in ACCOUNT_LIKE.findall(text):
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", m):
                continue
            if _norm(m) not in case_text:
                unprovided += 1
    return [("계좌 형식 값은 입력에서 온 것만 포함", unprovided == 0,
             "입력에 없는 계좌 형식 값이 발견됨" if unprovided else "")]


# ---------- 점검 6: 원고 ----------

def check_manuscript(case, out_dir, texts):
    files = [f for f in os.listdir(out_dir) if re.search(r"원고|강의안", f)]
    res = [("원고 파일을 만들지 않음(붙임에 이름만)", not files, _ids("field:attachments.manuscript"))]
    people = (case.get("allowance") or {}).get("people") or []
    uses = any((p.get("manuscript") or {}).get("kind", "none") != "none" for p in people)
    if uses:
        items, _ = parse_attachments(texts)
        res.append(("원고료 지급 시 붙임에 원고명이 있음", any(re.search("원고|강의안", n) for _, n in items),
                    _ids("field:draft.attachments")))
    return res


# ---------- 점검 7: HWPX ----------

def _hwpx_check_field(name):
    """Map checker labels to fixed schema paths; never echo labels from subprocess output."""
    normalized = (name or "").casefold()
    if "render_lines" in normalized:
        return "field:hwpx.render_text"
    if "mimetype" in normalized:
        return "field:hwpx.mimetype"
    if "xml" in normalized or "charpr" in normalized or "parapr" in normalized or "style" in normalized:
        return "field:hwpx.structure"
    if ("\ub4e4\uc5ec\uc4f0\uae30" in normalized
            or "\ub0b4\uc5b4\uc4f0\uae30" in normalized
            or "\uc904\ubc14\uafb8" in normalized):
        return "field:hwpx.paragraph_layout"
    if "\uacf5\ubc31" in normalized:
        return "field:hwpx.spacing"
    if "12pt" in normalized or "\uae00\uc790 \ud06c\uae30" in normalized:
        return "field:hwpx.font_size"
    if "\uae00\uaf34" in normalized:
        return "field:hwpx.font"
    return "field:hwpx.structure"


def _hwpx_failure_ids(stdout):
    """Keep fixed check IDs and checker-generated line numbers only."""
    failure_marker = "[\uC2E4\uD328]"
    safe_detail = re.compile(r"(?:line:\d+|paragraph-count)(?:, (?:line:\d+|paragraph-count))*")
    identifiers = []
    for line in (stdout or "").splitlines():
        if not line.startswith(failure_marker):
            continue
        tail = line[len(failure_marker):].strip()
        check_name, separator, detail = tail.partition(" - ")
        field = _hwpx_check_field(check_name.strip())
        identifiers.append(field)
        if separator and field in ("field:hwpx.paragraph_layout", "field:hwpx.render_text") and safe_detail.fullmatch(detail):
            identifiers.extend(f"field:hwpx.line[{number}]" for number in re.findall(r"line:(\d+)", detail))
    return _ids(*dict.fromkeys(("field:hwpx.structure", *identifiers)))


def check_hwpx(case, hwpx_path):
    with tempfile.TemporaryDirectory() as tmp:
        case_path = os.path.join(tmp, "case.json")
        with open(case_path, "w", encoding="utf-8") as f:
            json.dump(case, f, ensure_ascii=False)
        proc = subprocess.run([sys.executable, "-X", "utf8", "-I", os.path.join(BASE_DIR, "hwpx_check.py"),
                               hwpx_path, case_path], capture_output=True, text=True, encoding="utf-8")
    detail = _hwpx_failure_ids(proc.stdout) if proc.returncode else _ids("field:hwpx.structure")
    return [("HWPX 구조·본문·서식(hwpx_check.py)", proc.returncode == 0, detail)]


# ---------- 전체 ----------

def run_checks(case, out_dir):
    """결과 목록 [(항목, 통과 여부, 상세)]. case는 금액 채우기 전·후 어느 쪽이든 된다."""
    case = fill_draft_amounts(case)
    config = load_config()
    res = []
    hwpx = os.path.join(out_dir, HWPX_NAME)
    if not os.path.exists(hwpx):
        return [("기안문.hwpx 존재", False, _ids("field:hwpx.file"))]
    with zipfile.ZipFile(hwpx) as zf:
        texts = [p.lstrip() for p in extract_paragraphs(zf) if p.strip()]  # 줄 앞 들여쓰기 공백은 hwpx_check가 점검

    def guarded(name, fn):
        try:
            res.extend(fn())
        except Exception:  # 사용자 값이나 로컬 경로를 보고서에 남기지 않는다.
            res.append((name, False, _ids("field:check.execution")))

    guarded("(1) 금액", lambda: check_amounts(case, out_dir, texts))
    guarded("(2) 붙임", lambda: check_attachments(case, texts, config))
    guarded("(3) 항목기호·날짜", lambda: check_symbols_dates(case, texts))
    if case.get("allowance"):
        guarded("(4) 수당 엑셀", lambda: check_allowance_xlsx(case, out_dir))
        guarded("(4) 수당 단가 재계산", lambda: check_allowance_rates(case, out_dir))
    if case.get("purchase"):
        guarded("(4) 물품 엑셀", lambda: check_purchase_xlsx(case, out_dir))
    guarded("(4) 예산연결 업로드", lambda: check_budget_upload(case, out_dir))
    if case.get("expense"):
        guarded("(4) 경비 산출식", lambda: check_expense(case, texts))
    guarded("(5) 개인정보", lambda: check_privacy(case, out_dir))
    guarded("(6) 원고", lambda: check_manuscript(case, out_dir, texts))
    guarded("(7) HWPX", lambda: check_hwpx(case, hwpx))
    return res


def write_report(case, results, path):
    failed = [r for r in results if not r[1]]
    lines = ["# 자동 점검 결과", "",
             f"- 결과: {'전 항목 통과' if not failed else f'실패 {len(failed)}건'} (통과 {len(results) - len(failed)} / 전체 {len(results)})",
             "- 이 점검은 '자동 점검 통과'일 뿐, 서류가 다 갖춰진 '제출 준비 완료'와는 다릅니다(준비점검표 확인).", "",
             "| 결과 | 항목 | 상세 |", "|---|---|---|"]
    for name, ok, _detail in results:
        detail = "자동 점검" if ok else _safe_report_detail(_detail, name)
        lines.append(f"| {'통과' if ok else '**실패**'} | {name} | {detail} |")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    return failed


def main(argv):
    if len(argv) != 3:
        print("사용법: python harness.py case.json out_dir")
        return 2
    with open(argv[1], encoding="utf-8") as f:
        case = json.load(f)
    results = run_checks(case, argv[2])
    failed = write_report(case, results, os.path.join(argv[2], "점검결과.md"))
    for name, ok, detail in results:
        message = "" if ok else f" — {_safe_report_detail(detail, name)}"
        print(f"[{'통과' if ok else '실패'}] {name}{message}")
    print(f"\n통과 {len(results) - len(failed)} / 실패 {len(failed)}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
